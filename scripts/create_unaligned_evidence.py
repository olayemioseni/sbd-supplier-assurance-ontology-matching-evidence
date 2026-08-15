from pathlib import Path
from textwrap import fill

import pandas as pd
import matplotlib.pyplot as plt
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill
)
from openpyxl.utils import get_column_letter

PROJECT = Path(
    "/mnt/c/OntoMatch/sbd-supplier-assurance-evaluation"
)

OUTPUT_DIR = (
    PROJECT /
    "evaluation/consolidated_comparison/"
    "appendix_e_exports/unaligned_evidence"
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)

INPUTS = {
    "LogMap": (
        PROJECT /
        "evaluation/logmap_unaligned_review/"
        "logmap_unaligned_reciprocal_all_classified.csv"
    ),
    "AML": (
        PROJECT /
        "evaluation/aml_unaligned_review/"
        "aml_unaligned_reciprocal_all_classified.csv"
    ),
    "BERTMap": (
        PROJECT /
        "reproduction_outputs/BERTMap/evaluation/"
        "bertmap_reciprocal_candidates_classified.csv"
    )
}

SCOPES = {
    "LogMap": (
        "Reciprocal candidates across classes, "
        "object properties, data properties and individuals"
    ),
    "AML": (
        "Reciprocal candidates across classes, "
        "object properties, data properties and individuals"
    ),
    "BERTMap": (
        "Priority reciprocal class candidates only"
    )
}

MATCHER_ORDER = {
    "LogMap": 1,
    "AML": 2,
    "BERTMap": 3
}

CLASSIFICATION_ORDER = [
    "Full",
    "Partial",
    "Weak",
    "Missing"
]

def normalise_columns(frame):
    frame = frame.copy()

    frame.columns = [
        str(column).strip().lower()
        .replace(" ", "_")
        .replace("-", "_")
        for column in frame.columns
    ]

    return frame

def select_column(frame, candidates, default=""):
    for candidate in candidates:
        if candidate in frame.columns:
            return frame[candidate]

    return pd.Series(
        [default] * len(frame),
        index=frame.index
    )

def load_matcher(matcher, path):
    assert path.exists(), (
        f"Missing {matcher} evidence: {path}"
    )

    raw = normalise_columns(
        pd.read_csv(path)
    )

    result = pd.DataFrame({
        "Matcher": matcher,
        "Evaluation Scope": SCOPES[matcher],
        "Entity Type": select_column(
            raw,
            ["entity_type"],
            "class" if matcher == "BERTMap" else ""
        ),
        "Source IRI": select_column(
            raw,
            ["source_iri"]
        ),
        "Source Entity": select_column(
            raw,
            ["source_label", "source_entity"]
        ),
        "Source Definition": select_column(
            raw,
            ["source_definition"]
        ),
        "Target IRI": select_column(
            raw,
            ["target_iri"]
        ),
        "Target Entity": select_column(
            raw,
            ["target_label", "target_entity"]
        ),
        "Target Definition": select_column(
            raw,
            ["target_definition"]
        ),
        "Candidate Rank": select_column(
            raw,
            ["candidate_rank"]
        ),
        "Similarity": pd.to_numeric(
            select_column(
                raw,
                [
                    "combined_similarity",
                    "similarity",
                    "score"
                ]
            ),
            errors="coerce"
        ),
        "Ranking Band": select_column(
            raw,
            ["ranking_band"]
        ),
        "Reciprocal Best": select_column(
            raw,
            ["reciprocal_best"],
            True
        ),
        "Assurance Classification": select_column(
            raw,
            [
                "manual_classification",
                "assurance_classification"
            ]
        ),
        "Review Basis": select_column(
            raw,
            [
                "review_basis",
                "classification_reason",
                "review_notes"
            ]
        )
    })

    result = result.reset_index(drop=True)

    result.insert(
        0,
        "Candidate Number",
        range(1, len(result) + 1)
    )

    result.insert(
        0,
        "Evidence ID",
        [
            f"{matcher.upper()}-U-{number:03d}"
            for number in range(1, len(result) + 1)
        ]
    )

    result["Similarity"] = (
        result["Similarity"].round(6)
    )

    return result

matcher_frames = {
    matcher: load_matcher(matcher, path)
    for matcher, path in INPUTS.items()
}

combined = pd.concat(
    matcher_frames.values(),
    ignore_index=True
)

combined["Matcher Order"] = (
    combined["Matcher"].map(MATCHER_ORDER)
)

combined = (
    combined.sort_values(
        [
            "Matcher Order",
            "Candidate Number"
        ]
    )
    .drop(columns="Matcher Order")
    .reset_index(drop=True)
)

assert len(combined) == 254, (
    f"Expected 254 candidates, found {len(combined)}"
)

assert (
    combined["Assurance Classification"]
    .isin(CLASSIFICATION_ORDER)
    .all()
), "An invalid or missing classification was found"

# -------------------------------------------------
# Summary by matcher
# -------------------------------------------------
summary = pd.crosstab(
    combined["Matcher"],
    combined["Assurance Classification"]
)

summary = summary.reindex(
    index=["LogMap", "AML", "BERTMap"],
    columns=CLASSIFICATION_ORDER,
    fill_value=0
)

summary["Total"] = summary.sum(axis=1)

summary["Full or Partial"] = (
    summary["Full"] + summary["Partial"]
)

summary["Full/Partial Percentage"] = (
    summary["Full or Partial"]
    / summary["Total"]
)

summary = summary.reset_index()

summary.insert(
    1,
    "Evaluation Scope",
    summary["Matcher"].map(SCOPES)
)

# -------------------------------------------------
# Summary by matcher and entity type
# -------------------------------------------------
by_type = pd.crosstab(
    [
        combined["Matcher"],
        combined["Entity Type"]
    ],
    combined["Assurance Classification"]
)

by_type = by_type.reindex(
    columns=CLASSIFICATION_ORDER,
    fill_value=0
)

by_type["Total"] = by_type.sum(axis=1)
by_type = by_type.reset_index()

# -------------------------------------------------
# Method notes
# -------------------------------------------------
method_notes = pd.DataFrame({
    "Item": [
        "Purpose",
        "Candidate generation",
        "Reciprocal-best rule",
        "Full",
        "Partial",
        "Weak",
        "Missing",
        "Important limitation"
    ],
    "Explanation": [
        (
            "Supplementary examination of entities absent "
            "from the generated matcher alignments."
        ),
        (
            "TF-IDF label and definition similarity generated "
            "candidate target entities for unaligned SbD entities."
        ),
        (
            "A candidate was shortlisted where the source and "
            "target were each other's highest-ranked suggestion."
        ),
        (
            "The source and target represent semantically "
            "equivalent assurance concepts."
        ),
        (
            "The concepts meaningfully overlap but differ "
            "in scope, specificity or modelling role."
        ),
        (
            "Only limited assurance relevance exists."
        ),
        (
            "The candidate relationship is unsupported."
        ),
        (
            "These candidates are not matcher-generated "
            "mappings and must not be interpreted as recall."
        )
    ]
})

# -------------------------------------------------
# Save CSV files
# -------------------------------------------------
combined_csv = (
    OUTPUT_DIR /
    "all_unaligned_reciprocal_candidates_classified.csv"
)

summary_csv = (
    OUTPUT_DIR /
    "unaligned_candidate_summary_by_matcher.csv"
)

type_csv = (
    OUTPUT_DIR /
    "unaligned_candidate_summary_by_entity_type.csv"
)

combined.to_csv(
    combined_csv,
    index=False
)

summary.to_csv(
    summary_csv,
    index=False
)

by_type.to_csv(
    type_csv,
    index=False
)

# -------------------------------------------------
# Create Excel workbook
# -------------------------------------------------
workbook_file = (
    OUTPUT_DIR /
    "SbD_SORA_Unaligned_Candidate_Evidence.xlsx"
)

with pd.ExcelWriter(
    workbook_file,
    engine="openpyxl"
) as writer:

    summary.to_excel(
        writer,
        sheet_name="Matcher Summary",
        index=False
    )

    by_type.to_excel(
        writer,
        sheet_name="Summary by Entity Type",
        index=False
    )

    combined.to_excel(
        writer,
        sheet_name="All 254 Candidates",
        index=False
    )

    for matcher, frame in matcher_frames.items():
        frame.to_excel(
            writer,
            sheet_name=f"{matcher} Candidates",
            index=False
        )

    method_notes.to_excel(
        writer,
        sheet_name="Method Notes",
        index=False
    )

    header_fill = PatternFill(
        "solid",
        fgColor="203864"
    )

    classification_colours = {
        "Full": "C6E0B4",
        "Partial": "FFE699",
        "Weak": "F4B183",
        "Missing": "F4CCCC"
    }

    for sheet in writer.sheets.values():
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(
                color="FFFFFF",
                bold=True
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        for column_cells in sheet.columns:
            letter = get_column_letter(
                column_cells[0].column
            )

            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells[:200]
            )

            sheet.column_dimensions[letter].width = min(
                max(max_length + 2, 12),
                55
            )

        classification_column = None

        for cell in sheet[1]:
            if cell.value == "Assurance Classification":
                classification_column = cell.column
                break

        if classification_column:
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(
                    row=row,
                    column=classification_column
                )

                colour = classification_colours.get(
                    str(cell.value)
                )

                if colour:
                    cell.fill = PatternFill(
                        "solid",
                        fgColor=colour
                    )

# -------------------------------------------------
# Create summary picture
# -------------------------------------------------
summary_picture = summary.copy()

summary_picture["Full/Partial Percentage"] = (
    summary_picture["Full/Partial Percentage"]
    .map(lambda value: f"{value:.2%}")
)

fig, ax = plt.subplots(
    figsize=(18, 4.8)
)

ax.axis("off")

ax.set_title(
    "Supplementary Unaligned-Candidate Classification",
    fontsize=17,
    fontweight="bold",
    color="#203864",
    pad=18
)

summary_display = summary_picture[
    [
        "Matcher",
        "Full",
        "Partial",
        "Weak",
        "Missing",
        "Total",
        "Full or Partial",
        "Full/Partial Percentage"
    ]
]

table = ax.table(
    cellText=summary_display.values,
    colLabels=summary_display.columns,
    cellLoc="center",
    colLoc="center",
    loc="center"
)

table.auto_set_font_size(False)
table.set_fontsize(10)
table.scale(1, 1.8)

for (row, column), cell in table.get_celld().items():
    cell.set_edgecolor("#A6A6A6")

    if row == 0:
        cell.set_facecolor("#203864")
        cell.set_text_props(
            color="white",
            weight="bold"
        )
    elif row % 2 == 0:
        cell.set_facecolor("#D9EAF7")
    else:
        cell.set_facecolor("white")

summary_png = (
    OUTPUT_DIR /
    "unaligned_candidate_summary.png"
)

plt.savefig(
    summary_png,
    dpi=220,
    bbox_inches="tight",
    facecolor="white"
)

plt.close(fig)

# -------------------------------------------------
# Create full picture for each matcher
# -------------------------------------------------
def create_matcher_picture(matcher, frame):
    picture = frame[
        [
            "Candidate Number",
            "Entity Type",
            "Source Entity",
            "Target Entity",
            "Similarity",
            "Assurance Classification"
        ]
    ].copy()

    picture["Source Entity"] = (
        picture["Source Entity"]
        .apply(lambda value: fill(str(value), 34))
    )

    picture["Target Entity"] = (
        picture["Target Entity"]
        .apply(lambda value: fill(str(value), 34))
    )

    picture["Similarity"] = (
        picture["Similarity"]
        .map(
            lambda value: (
                f"{value:.4f}"
                if pd.notna(value)
                else ""
            )
        )
    )

    height = max(
        10,
        0.39 * (len(picture) + 3)
    )

    fig, ax = plt.subplots(
        figsize=(20, height)
    )

    ax.axis("off")

    ax.set_title(
        f"{matcher} Supplementary Unaligned Candidates",
        fontsize=17,
        fontweight="bold",
        color="#203864",
        pad=18
    )

    table = ax.table(
        cellText=picture.values,
        colLabels=picture.columns,
        cellLoc="left",
        colLoc="left",
        loc="upper center",
        colWidths=[
            0.08,
            0.12,
            0.27,
            0.27,
            0.10,
            0.16
        ]
    )

    table.auto_set_font_size(False)
    table.set_fontsize(7)
    table.scale(1, 1.4)

    classification_column = list(
        picture.columns
    ).index("Assurance Classification")

    colours = {
        "Full": "#C6E0B4",
        "Partial": "#FFE699",
        "Weak": "#F4B183",
        "Missing": "#F4CCCC"
    }

    for (row, column), cell in table.get_celld().items():
        cell.set_edgecolor("#A6A6A6")
        cell.set_linewidth(0.5)

        if row == 0:
            cell.set_facecolor("#203864")
            cell.set_text_props(
                color="white",
                weight="bold"
            )
        elif row % 2 == 0:
            cell.set_facecolor("#D9EAF7")
        else:
            cell.set_facecolor("white")

        if row > 0 and column == classification_column:
            category = picture.iloc[
                row - 1
            ]["Assurance Classification"]

            cell.set_facecolor(
                colours.get(category, "white")
            )

    output_file = (
        OUTPUT_DIR /
        f"{matcher.lower()}_unaligned_candidates.png"
    )

    plt.savefig(
        output_file,
        dpi=200,
        bbox_inches="tight",
        facecolor="white"
    )

    plt.close(fig)

    return output_file

matcher_pictures = [
    create_matcher_picture(matcher, frame)
    for matcher, frame in matcher_frames.items()
]

print("UNALIGNED EVIDENCE CREATED")
print("==========================")
print(summary.to_string(index=False))

print()
print("Total candidates:", len(combined))
print("Workbook:", workbook_file)
print("Combined CSV:", combined_csv)
print("Summary CSV:", summary_csv)
print("Entity-type CSV:", type_csv)
print("Summary image:", summary_png)

for picture in matcher_pictures:
    print("Matcher image:", picture)
